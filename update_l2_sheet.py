#!/usr/bin/env python3
"""
Pushes the L2 discount map's data into the Google Sheet this project's Apps
Script reads from, as two read-only reference tabs ("L2 Points", "Condition")
-- so the discount targeting data lives in the Sheet, never in this repo.

Run this whenever the source Project Atlas xlsx (or the village file it's
joined against) changes. Each run fully replaces both tabs (not an
incremental append), so it's always safe to re-run.

index.html has no L2 data baked into it -- it fetches everything live from
the "L2 Points" / "Condition" tabs this script writes to, after the viewer
signs in and the Apps Script backend checks their email against the Users
tab. So running this script is the whole update: nothing about the page
needs rebuilding.

What this reproduces (originally worked out ad hoc against the source
workbook, see this project's README for the fuller story):
  - The FULL "Condition" sheet -- all 6 archetypes x 3 ARPA bands x 2 port
    groups = 36 rows, each with its own 9 MKT-package sub-table -- not a
    shortcut limited to whichever combinations happen to appear in the
    current L2 data.
  - "L2 Data vAug" filtered down to every row whose (Archetype, ARPA group,
    Port group, NAD flag) resolves, via that same Condition table, to a
    discount greater than zero.
  - Each kept L2's phy_resre_nbr joined against Vill_SPLITTER_L2 in the
    Active FTTH village file for a human-readable village name (a clean 1:1
    match, no ambiguity -- confirmed by scanning the whole file -- but only
    around 45% of L2s here have an active-customer record to match against;
    the rest simply have no village name, not a wrong one).
"""
import csv
import json
import urllib.request
import urllib.error
from pathlib import Path

import openpyxl

# Keep this in sync with DEFAULT_SYNC_URL in index.html -- if you change one,
# change the other.
SYNC_URL = "REPLACE_WITH_YOUR_APPS_SCRIPT_WEB_APP_URL"

DASHBOARD_DIR = Path(__file__).resolve().parent.parent


def _find(name, folders):
    """Locate a file that lives outside this repo, in the Dashboard folder.
    Tries each candidate folder in order, falls back to the first one in the
    error message so it names a sensible location to create the file in."""
    for folder in folders:
        candidate = DASHBOARD_DIR / folder / name
        if candidate.exists():
            return candidate
    return DASHBOARD_DIR / folders[0] / name


L2_XLSX = _find("20260801 Project Atlas L2 Aug'26 vNTB Pak Kret.xlsx", ["L2", "Config", "."])
VILLAGE_FILE = _find("Active FTTH In Village_BMA-West.TXT", [Path("TOL") / "Data", "Config", "."])
SYNC_SECRET_FILE = DASHBOARD_DIR / "Config" / "l2_sync_secret.txt"


def load_village_lookup(path):
    """Vill_SPLITTER_L2 -> GIS_VILL_NAME, first match wins (verified 1:1 --
    no L2 in the source file ever maps to more than one village name)."""
    lookup = {}
    with open(path, encoding="utf-8-sig") as f:
        reader = csv.reader(f, delimiter="^")
        header = next(reader)
        idx_l2 = header.index("Vill_SPLITTER_L2")
        idx_vname = header.index("GIS_VILL_NAME")
        for row in reader:
            if len(row) <= max(idx_l2, idx_vname):
                continue
            l2, vname = row[idx_l2].strip(), row[idx_vname].strip()
            if l2 and vname and l2 not in lookup:
                lookup[l2] = vname
    return lookup


def load_conditions(wb):
    """Every row of the Condition sheet, verbatim -- all 36 Archetype x
    ARPA-band x Port-group combinations (NAD flag is always 00_Pass in this
    table; a row with NAD flagged for removal simply has no entry here)."""
    ws = wb["Condition"]
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    conditions = []
    for r in rows:
        if r[1] is None:
            continue
        arch, arpa_group, port, nad, disc_pct = r[1], r[2], r[3], r[4], r[5]
        mkts = []
        idx = 6
        for _ in range(9):
            code, desc, disc, normal, special = r[idx], r[idx + 1], r[idx + 2], r[idx + 3], r[idx + 4]
            mkts.append({
                "code": code, "desc": desc,
                "disc": round(disc, 4), "normal": normal, "special": round(special, 2),
            })
            idx += 5
        conditions.append({
            "arch": arch, "arpaGroup": arpa_group, "port": port, "nad": nad,
            "discPct": round(disc_pct, 4), "mkts": mkts,
        })
    return conditions


def find_condition(conditions, arch, arpa_group, port, nad):
    for c in conditions:
        if c["arch"] == arch and c["arpaGroup"] == arpa_group and c["port"] == port and c["nad"] == nad:
            return c
    return None


def load_points(wb, conditions, village_lookup):
    """Every L2 Data row whose Condition-table lookup resolves to a nonzero
    discount -- a genuine per-row lookup against the full table, not a
    shortcut for whichever combinations happen to appear in this file."""
    ws = wb["L2 Data vAug"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    points = []
    for r in rows:
        arch, arpa_group, port, nad = r[4], r[13], r[12], r[14]
        cond = find_condition(conditions, arch, arpa_group, port, nad)
        if cond is None or cond["discPct"] <= 0:
            continue
        l2_id = r[10]  # phy_resre_nbr
        points.append({
            "id": l2_id,
            "hpb": r[2],
            "lat": round(r[8], 6),
            "lon": round(r[9], 6),
            "arch": arch,
            "port": port,
            "arpa": round(r[11]),
            "arpaGroup": arpa_group,
            "nad": nad,
            "adm2": r[6],
            "adm3": r[7],
            "village": village_lookup.get(l2_id),
        })
    return points


def main():
    if not L2_XLSX.exists():
        raise SystemExit(f"Not found: {L2_XLSX}")
    if not VILLAGE_FILE.exists():
        raise SystemExit(f"Not found: {VILLAGE_FILE}")
    if not SYNC_SECRET_FILE.exists():
        raise SystemExit(
            f"Not found: {SYNC_SECRET_FILE}\n"
            "Create it containing the same value as the SYNC_SECRET Script "
            "Property in the Apps Script project, with no extra whitespace."
        )
    sync_secret = SYNC_SECRET_FILE.read_text(encoding="utf-8").strip()

    if SYNC_URL.startswith("REPLACE_WITH"):
        raise SystemExit(
            "SYNC_URL at the top of this script still needs to be set to "
            "your Apps Script Web App URL -- see the repo README."
        )

    print("Reading village name lookup...")
    village_lookup = load_village_lookup(VILLAGE_FILE)
    print(f"  {len(village_lookup)} L2 splitters have a village name on file")

    print("Reading source workbook...")
    wb = openpyxl.load_workbook(L2_XLSX, data_only=True, read_only=True)
    conditions = load_conditions(wb)
    points = load_points(wb, conditions, village_lookup)
    matched = sum(1 for p in points if p["village"])
    print(f"Parsed {len(conditions)} condition rows, {len(points)} eligible L2 points "
          f"({matched} with a matched village name)")

    payload = json.dumps({
        "action": "syncL2Data", "points": points, "conditions": conditions, "secret": sync_secret,
    }).encode("utf-8")
    req = urllib.request.Request(
        SYNC_URL, data=payload, method="POST",
        headers={"Content-Type": "text/plain;charset=utf-8"},
    )
    print("Uploading to Google Sheet...")
    try:
        with urllib.request.urlopen(req, timeout=60) as res:
            body = res.read().decode("utf-8")
    except urllib.error.HTTPError as e:
        raise SystemExit(f"Upload failed: HTTP {e.code}\n{e.read().decode('utf-8', 'replace')[:500]}")
    except urllib.error.URLError as e:
        raise SystemExit(f"Upload failed: {e.reason}")

    try:
        result = json.loads(body)
    except json.JSONDecodeError:
        raise SystemExit(
            "Upload failed: response wasn't JSON (the Web App URL may need to be "
            "redeployed with \"Who has access: Anyone\").\n"
            f"First 300 chars of response:\n{body[:300]}"
        )

    if result.get("ok"):
        print(f"Done -- {result.get('pointCount', len(points))} points, "
              f"{result.get('conditionRowCount', '?')} condition rows synced to the sheet.")
    else:
        raise SystemExit(f"Upload failed: {result.get('error')}")


if __name__ == "__main__":
    main()
